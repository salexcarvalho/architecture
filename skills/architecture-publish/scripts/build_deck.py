#!/usr/bin/env python3
"""
Gerador reutilizável de decks de arquitetura resumidos (.pptx), 16:9.

Uso:  build_deck.py CONTEUDO.json [SAIDA.pptx] [--render DIR]
      SAIDA padrão: campo "output" do JSON, ou CONTEUDO.pptx ao lado do JSON.
      --render DIR: converte o .pptx em PDF (soffice) e em PNG por slide (pdftoppm)
                    dentro de DIR, para conferência visual.
Dep.: python-pptx (use um venv; nunca instalação global). Render: libreoffice + poppler.

O conteúdo fica no JSON (ver templates/deck-conteudo.example.json); este arquivo só
sabe desenhar. Tipos de slide:
  cover    {topo, titulo, subtitulo, meta}
  divider  {kicker, titulo}
  bullets  {kicker, titulo, itens:[[texto, sub|null]], nota?}
  table    {kicker, titulo, cols, rows, widths?, fs?, altura?, nota?}
  cards    {kicker, titulo, itens:[[rotulo, texto, tom]], cols?}
  callout  {kicker, titulo, frase, corpo:[...], tom?}
  flow     {kicker, titulo, passos:[[nome, legenda]], cards:[[rotulo, texto, tom]]}
  diagram  {kicker, titulo, boxes:[...], arrows:[...], textos?:[...]}
  closing  {titulo, itens:[texto | [texto, detalhe]], rodape}
tom: accent | ok | warn | crit | plain
Opcional no JSON: "theme": {"ACCENT": "0E5A55", "DARK": "0E1E1C", ...}
Regras de conteúdo: título = afirmação; sem IDs internos, custos nem nomes de pessoas.
"""
import argparse
import json
import pathlib
import shutil
import subprocess
import sys


def _cli():
    ap = argparse.ArgumentParser(
        description="Gera o deck resumido de arquitetura (.pptx) a partir de um JSON de conteúdo.",
        epilog="Dependência: python-pptx (venv). Exemplo: build_deck.py deck.json Sistema-Arquitetura-Resumo.pptx --render /tmp/render")
    ap.add_argument("conteudo", help="JSON de conteúdo (ver deck-conteudo.example.json)")
    ap.add_argument("saida", nargs="?", default=None, help="arquivo .pptx de saída")
    ap.add_argument("--render", metavar="DIR", default=None,
                    help="gera PDF e PNG por slide em DIR (soffice + pdftoppm) para conferência visual")
    return ap.parse_args()


ARGS = _cli() if __name__ == "__main__" else None  # --help não exige python-pptx

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

HERE = pathlib.Path(__file__).parent

# paleta e tipografia do padrao Montreal
INK       = RGBColor(0x0F, 0x18, 0x17)
INK2      = RGBColor(0x38, 0x4A, 0x47)
MUTED     = RGBColor(0x61, 0x73, 0x6F)
ACCENT    = RGBColor(0x0E, 0x5A, 0x55)
ACCENT_LT = RGBColor(0x5F, 0xC9, 0xB8)
SOFT      = RGBColor(0xDB, 0xEB, 0xE8)
GROUND    = RGBColor(0xF3, 0xF6, 0xF5)
SURFACE   = RGBColor(0xFF, 0xFF, 0xFF)
LINE      = RGBColor(0xCF, 0xDA, 0xD7)
DARK      = RGBColor(0x0E, 0x1E, 0x1C)
OK        = RGBColor(0x1B, 0x6E, 0x45)
OK_S      = RGBColor(0xDC, 0xED, 0xE4)
WARN      = RGBColor(0x87, 0x59, 0x0A)
WARN_S    = RGBColor(0xF4, 0xE7, 0xCB)
CRIT      = RGBColor(0x9A, 0x2C, 0x25)
CRIT_S    = RGBColor(0xF4, 0xDB, 0xD9)
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
PALE      = RGBColor(0xC9, 0xD9, 0xD5)   # texto claro sobre fundo escuro
DIM       = RGBColor(0x7C, 0x91, 0x8C)   # texto de apoio sobre fundo escuro

FONT = "Calibri"
W, H = 13.333, 7.5


def apply_theme(theme):
    g = globals()
    for nome, hexa in (theme or {}).items():
        if nome in g and isinstance(g[nome], RGBColor):
            g[nome] = RGBColor.from_string(hexa.lstrip("#").upper())


def new_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    return prs


def _blank(prs, bg=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bgfill = s.background.fill
    bgfill.solid()
    bgfill.fore_color.rgb = bg if bg is not None else GROUND
    return s


def _txt(slide, x, y, w, h, text, size=18, bold=False, color=INK,
         align=PP_ALIGN.LEFT, space=6, italic=False, line=1.25,
         anchor=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if anchor is not None:
        tf.vertical_anchor = anchor
    lines = text if isinstance(text, list) else [text]
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(space)
        p.line_spacing = line
        run = p.add_run()
        run.text = ln
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = FONT
    return box


def _rect(slide, x, y, w, h, fill, line_col=None):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = 0.06
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line_col:
        sh.line.color.rgb = line_col
        sh.line.width = Pt(1)
    else:
        sh.line.fill.background()
    sh.shadow.inherit = False
    sh.text_frame.text = ""
    return sh


def _tons():
    return {"accent": (SOFT, ACCENT), "ok": (OK_S, OK),
            "warn": (WARN_S, WARN), "crit": (CRIT_S, CRIT),
            "plain": (SURFACE, INK2)}


def _arrow(slide, x1, y1, x2, y2, color=None, width=1.75):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color if color is not None else ACCENT
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return c


# ───────────────────────────────────────────── tipos de slide
def cover(prs, topo, titulo, subtitulo, meta):
    s = _blank(prs, DARK)
    _rect(s, 0, 6.95, W, 0.55, ACCENT)
    _txt(s, 1.0, 1.7, 11, 0.5, str(topo).upper(), size=15, bold=True, color=ACCENT_LT)
    _txt(s, 1.0, 2.35, 11, 2.0, titulo, size=42, bold=True, color=WHITE, line=1.1)
    _txt(s, 1.0, 4.45, 10.5, 1.2, subtitulo, size=19, color=PALE, line=1.35)
    _txt(s, 1.0, 6.05, 11, 0.7, meta, size=13, color=DIM)
    return s


def divider(prs, kicker, titulo):
    s = _blank(prs, DARK)
    _rect(s, 0, 3.05, 0.55, 1.4, ACCENT_LT)
    _txt(s, 1.0, 3.05, 11, 0.4, kicker, size=14, bold=True, color=ACCENT_LT)
    _txt(s, 1.0, 3.55, 11, 1.2, titulo, size=34, bold=True, color=WHITE, line=1.15)
    return s


def head(slide, kicker, titulo):
    # titulo longo quebra em duas linhas: reduz o corpo para nao colidir com a regua
    size = 30 if len(titulo) <= 46 else 25
    _txt(slide, 0.9, 0.55, 11.5, 0.35, kicker.upper(), size=13, bold=True, color=ACCENT)
    _txt(slide, 0.9, 0.95, 11.5, 0.95, titulo, size=size, bold=True, color=INK, line=1.1)
    ln = slide.shapes.add_shape(1, Inches(0.9), Inches(1.95), Inches(11.5), Emu(9525))
    ln.fill.solid(); ln.fill.fore_color.rgb = LINE; ln.line.fill.background()
    ln.shadow.inherit = False


def bullets(prs, kicker, titulo, itens, nota=None):
    s = _blank(prs)
    head(s, kicker, titulo)
    base = 6.15 if nota else 6.95
    pesos = [1 + (len(sub) // 105 if sub else -0.35) for _, sub in itens]
    unidade = min(0.92, (base - 2.2) / max(sum(pesos), 1))
    apertado = unidade < 0.85
    fs_t, fs_s = (15, 12.5) if apertado else (17, 14)
    y = 2.2
    for (texto, sub), peso in zip(itens, pesos):
        _rect(s, 0.9, y + 0.11, 0.13, 0.13, ACCENT)
        _txt(s, 1.25, y, 11.1, 0.4, texto, size=fs_t, bold=True, color=INK)
        if sub:
            _txt(s, 1.25, y + (0.35 if apertado else 0.42), 11.0, 0.5, sub,
                 size=fs_s, color=INK2, line=1.3)
        y += unidade * peso
    if nota:
        _rect(s, 0.9, 6.3, 11.5, 0.75, SOFT)
        _txt(s, 1.15, 6.48, 11.0, 0.5, nota, size=14, bold=True, color=ACCENT)
    return s


def table(prs, kicker, titulo, cols, rows, widths=None, nota=None, fs=12.5, altura=None):
    s = _blank(prs)
    head(s, kicker, titulo)
    n_r, n_c = len(rows) + 1, len(cols)
    top = 2.2
    height = altura or min(4.0, 0.42 + 0.44 * len(rows))
    shape = s.shapes.add_table(n_r, n_c, Inches(0.9), Inches(top),
                               Inches(11.5), Inches(height))
    tbl = shape.table
    if widths:
        total = sum(widths)
        for i, wd in enumerate(widths):
            tbl.columns[i].width = Emu(int(Inches(11.5) * wd / total))

    def _cell(cell, texto, size, bold, color, bg):
        cell.fill.solid(); cell.fill.fore_color.rgb = bg
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.margin_left = cell.margin_right = Inches(0.09)
        tf = cell.text_frame
        tf.word_wrap = True
        for k, ln in enumerate(str(texto).split("\n")):
            p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
            p.space_after = Pt(0)
            run = p.add_run()
            run.text = ln
            run.font.size = Pt(size); run.font.bold = bold
            run.font.color.rgb = color; run.font.name = FONT

    for j, c in enumerate(cols):
        _cell(tbl.cell(0, j), c, fs, True, WHITE, ACCENT)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            _cell(tbl.cell(i, j), val, fs, j == 0, INK2,
                  SURFACE if i % 2 else GROUND)
    if nota:
        _txt(s, 0.9, top + height + 0.25, 11.5, 0.8, nota, size=14,
             color=ACCENT, bold=True, line=1.3)
    return s


def callout(prs, kicker, titulo, frase, corpo, tom="accent"):
    bg, fg = _tons()[tom]
    s = _blank(prs)
    head(s, kicker, titulo)
    _rect(s, 0.9, 2.25, 11.5, 2.15, bg)
    _rect(s, 0.9, 2.25, 0.09, 2.15, fg)
    fs_frase = 23 if len(frase) <= 174 else (20 if len(frase) <= 250 else 18)
    _txt(s, 1.35, 2.6, 10.6, 1.6, frase, size=fs_frase, bold=True, color=fg, line=1.22)
    y = 4.75
    for linha in corpo:
        _rect(s, 0.95, y + 0.1, 0.11, 0.11, fg)
        _txt(s, 1.3, y, 11.0, 0.55, linha, size=15, color=INK2, line=1.3)
        y += 0.52 + 0.20 * (len(linha) // 110)
    return s


def _card(s, x, y, cw, ch, rot, val, tom, fs_val=14):
    bg, fg = _tons()[tom]
    _rect(s, x, y, cw, ch, bg, LINE if tom == "plain" else None)
    _txt(s, x + 0.3, y + 0.22, cw - 0.6, 0.35, rot, size=12.5, bold=True, color=fg)
    _txt(s, x + 0.3, y + 0.62, cw - 0.6, ch - 0.75, val, size=fs_val, color=INK2, line=1.26)


def cards(prs, kicker, titulo, itens, cols=None):
    s = _blank(prs)
    head(s, kicker, titulo)
    n = len(itens)
    cols = cols or (3 if n > 4 else 2)
    cw = (11.5 - 0.35 * (cols - 1)) / cols
    ch = 2.0
    for i, (rot, val, tom) in enumerate(itens):
        r, c = divmod(i, cols)
        _card(s, 0.9 + c * (cw + 0.35), 2.25 + r * (ch + 0.3), cw, ch, rot, val, tom)
    return s


def flow(prs, kicker, titulo, passos, cards_=None):
    """Sequencia de estados (caixas ligadas por setas) + cards de apoio."""
    s = _blank(prs)
    head(s, kicker, titulo)
    n = len(passos)
    gap = 0.34
    bw = (11.5 - gap * (n - 1)) / n
    bh, y0 = 1.3, 2.3
    for i, (nome, legenda) in enumerate(passos):
        ultimo = i == n - 1
        bg, fg = (OK_S, OK) if ultimo else (SOFT, ACCENT)
        x = 0.9 + i * (bw + gap)
        _rect(s, x, y0, bw, bh, bg)
        _txt(s, x + 0.12, y0 + 0.2, bw - 0.24, 0.35, f"{i + 1} · {nome}", size=14.5,
             bold=True, color=fg, align=PP_ALIGN.CENTER)
        _txt(s, x + 0.12, y0 + 0.62, bw - 0.24, 0.6, legenda, size=12, color=INK2,
             align=PP_ALIGN.CENTER, line=1.2)
        if not ultimo:
            _arrow(s, x + bw + 0.04, y0 + bh / 2, x + bw + gap - 0.04, y0 + bh / 2)
    if cards_:
        k = len(cards_)
        cw = (11.5 - 0.35 * (k - 1)) / k
        for i, (rot, val, tom) in enumerate(cards_):
            _card(s, 0.9 + i * (cw + 0.35), 4.1, cw, 2.35, rot, val, tom)
    return s


def diagram(prs, kicker, titulo, boxes, arrows=None, textos=None):
    """boxes: {id,x,y,w,h,titulo,sub,tom,grupo?}; grupo = moldura com rotulo.
    arrows: {x1,y1,x2,y2,label?,lx,ly,lw}. textos: {x,y,w,h,texto,size}."""
    s = _blank(prs)
    head(s, kicker, titulo)
    tons = _tons()
    for b in boxes:  # molduras primeiro, para ficarem atras
        if "grupo" in b:
            _rect(s, b["x"], b["y"], b["w"], b["h"], SURFACE, LINE)
            _txt(s, b["x"] + 0.2, b["y"] + 0.15, b["w"] - 0.4, 0.3, b["grupo"],
                 size=11.5, bold=True, color=MUTED)
    for b in boxes:
        if "grupo" in b:
            continue
        bg, fg = tons[b.get("tom", "accent")]
        _rect(s, b["x"], b["y"], b["w"], b["h"], bg)
        sub = b.get("sub")
        tam = 14 if b["h"] >= 0.9 else 13.5
        if sub:
            _txt(s, b["x"] + 0.1, b["y"] + 0.14, b["w"] - 0.2, 0.3, b["titulo"],
                 size=tam, bold=True, color=fg, align=PP_ALIGN.CENTER)
            _txt(s, b["x"] + 0.1, b["y"] + 0.14 + 0.32 + (0.05 if b["h"] >= 0.9 else 0),
                 b["w"] - 0.2, b["h"] - 0.5, sub, size=11.5, color=INK2,
                 align=PP_ALIGN.CENTER, line=1.15, space=0)
        else:
            _txt(s, b["x"] + 0.1, b["y"], b["w"] - 0.2, b["h"], b["titulo"], size=tam,
                 bold=True, color=fg, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    for a in arrows or []:
        _arrow(s, a["x1"], a["y1"], a["x2"], a["y2"])
        if a.get("label"):
            _txt(s, a["lx"], a["ly"], a.get("lw", 1.5), 0.25, a["label"], size=11,
                 italic=True, color=MUTED)
    for t in textos or []:
        _txt(s, t["x"], t["y"], t["w"], t.get("h", 1), t["texto"],
             size=t.get("size", 13), color=INK2, line=1.3)
    return s


def closing(prs, titulo, itens, rodape):
    """itens: texto ou [texto, detalhe]. Espacamento adaptativo."""
    s = _blank(prs, DARK)
    _txt(s, 1.0, 1.25, 11, 0.9, titulo, size=30, bold=True, color=WHITE)
    n = len(itens)
    topo, base = 2.35, 6.15
    passo = min(0.78, (base - topo) / max(n, 1))
    lado = min(0.42, passo - 0.12)
    fs = 17 if n <= 5 else 15
    y = topo
    for i, it in enumerate(itens, 1):
        texto, detalhe = (it, None) if isinstance(it, str) else (it[0], it[1])
        _rect(s, 1.0, y, lado, lado, ACCENT)
        _txt(s, 1.0, y + lado * 0.14, lado, lado * 0.8, str(i), size=fs - 2,
             bold=True, color=WHITE, align=PP_ALIGN.CENTER)
        x = 1.0 + lado + 0.23
        _txt(s, x, y + 0.0, 10.6, 0.35, texto, size=fs, bold=bool(detalhe),
             color=WHITE if detalhe else PALE, line=1.15)
        if detalhe:
            _txt(s, x, y + 0.33, 10.6, 0.3, detalhe, size=fs - 4, color=ACCENT_LT, line=1.15)
        y += passo
    _rect(s, 0, 6.95, W, 0.55, ACCENT)
    _txt(s, 1.0, 6.42, 11, 0.4, rodape, size=12, color=DIM)
    return s


BUILDERS = {
    "cover": cover, "divider": divider, "bullets": bullets, "table": table,
    "callout": callout, "cards": cards, "flow": flow, "diagram": diagram,
    "closing": closing,
}


def build(conteudo, saida):
    apply_theme(conteudo.get("theme"))
    prs = new_deck()
    for n, spec in enumerate(conteudo["slides"], 1):
        spec = dict(spec)
        tipo = spec.pop("type")
        if tipo == "flow" and "cards" in spec:
            spec["cards_"] = spec.pop("cards")
        if tipo == "bullets":
            spec["itens"] = [tuple(i) for i in spec["itens"]]
        if tipo in ("cards",):
            spec["itens"] = [tuple(i) for i in spec["itens"]]
        try:
            BUILDERS[tipo](prs, **spec)
        except Exception as e:  # mostra qual slide falhou
            raise SystemExit(f"slide {n} ({tipo}): {e!r}")
    prs.save(saida)
    return len(prs.slides)


def render(pptx, destino):
    """pptx -> PDF (soffice) -> PNG por slide (pdftoppm). Retorna a lista de PNGs."""
    destino = pathlib.Path(destino)
    destino.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise SystemExit("soffice/libreoffice não encontrado: renderização indisponível")
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(destino), str(pptx)],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    pdf = destino / (pathlib.Path(pptx).stem + ".pdf")
    if shutil.which("pdftoppm"):
        subprocess.run(["pdftoppm", "-r", "60", "-png", str(pdf), str(destino / "slide")], check=True)
    return sorted(destino.glob("slide*.png"))


def main():
    src = pathlib.Path(ARGS.conteudo)
    conteudo = json.loads(src.read_text(encoding="utf-8"))
    saida = pathlib.Path(ARGS.saida) if ARGS.saida else src.parent / conteudo.get(
        "output", src.with_suffix(".pptx").name)
    n = build(conteudo, saida)
    print(f"{n} slides -> {saida}")
    if ARGS.render:
        pngs = render(saida, ARGS.render)
        print(f"render: {len(pngs)} PNG em {ARGS.render}")


if __name__ == "__main__":
    main()
