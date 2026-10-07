#!/usr/bin/env python3
"""Monta o PPTX: cada slide é um PNG 16:9 em tela cheia (fiel ao PDF); o texto do slide vai nas notas.

Uso: build_pptx.py WORKDIR SAIDA.pptx   (requer python-pptx; use um venv)
Lê WORKDIR/assets/slides/slide-*.png, WORKDIR/assets/notes.json e WORKDIR/conteudo-deck.json (meta).
"""
import glob
import json
import os
import sys

from pptx import Presentation
from pptx.util import Emu

wd, out = os.path.abspath(sys.argv[1]), sys.argv[2]
pngs = sorted(glob.glob(os.path.join(wd, "assets", "slides", "slide-*.png")))
notes = json.load(open(os.path.join(wd, "assets", "notes.json"), encoding="utf-8"))
meta = json.load(open(os.path.join(wd, "conteudo-deck.json"), encoding="utf-8")).get("meta", {})
assert len(pngs) == len(notes), f"PNGs ({len(pngs)}) != slides ({len(notes)})"
prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)  # 16:9
blank = prs.slide_layouts[6]
for p, n in zip(pngs, notes):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(p, 0, 0, prs.slide_width, prs.slide_height)
    s.notes_slide.notes_text_frame.text = n
prs.core_properties.title = meta.get("titulo_html", meta.get("sistema", "Apresentação de arquitetura"))
prs.core_properties.author = meta.get("autor", "")
prs.save(out)
print("pptx:", out, len(pngs), "slides")
